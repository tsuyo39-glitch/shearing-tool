import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from core.exporter import export_result
from core.models import Placement, ProductSpec, SheetPlan, SheetType
from core.optimizer import group_sheet_plans, optimize, validate_layout
from core.weight import normalize_spec, plate_weight_kg


class WeightTests(unittest.TestCase):
    def test_z08_weight(self):
        weight, warning = plate_weight_kg("SGCC-Z08-NC", 1.2, 1000, 1000)
        self.assertAlmostEqual(weight, 9.5, places=6)
        self.assertIsNone(warning)

    def test_plain_steel_weight(self):
        weight, warning = plate_weight_kg("SPCC", 1.0, 1000, 2000)
        self.assertAlmostEqual(weight, 15.7, places=6)
        self.assertIsNone(warning)

    def test_nfkc_and_z8_normalization(self):
        self.assertEqual(normalize_spec(" ＳＧＣＣ－Ｚ８－ＮＣ "), "SGCC-Z8-NC")
        weight, warning = plate_weight_kg("SGCC-Z8-NC", 1.2, 1000, 1000)
        self.assertAlmostEqual(weight, 9.5, places=6)
        self.assertIsNone(warning)


class OptimizerTests(unittest.TestCase):
    def test_one_piece_uses_nominal_sheet_dimensions(self):
        product = ProductSpec("P1", "SPCC", 1.0, 980, 980)
        sheet = SheetType("S1", "大板", "SPCC", 1.0, 1000, 1000, edge_loss=10)
        result = optimize([product], [sheet], 1)[0]
        self.assertTrue(result.complete)
        self.assertAlmostEqual(result.yield_rate, 96.04, places=2)
        self.assertEqual(validate_layout(result), [])
        placement = result.sheets[0].placements[0]
        self.assertEqual((placement.x, placement.y), (0, 0))

    def test_edge_loss_does_not_reduce_fit_dimensions(self):
        product = ProductSpec("P1", "SPCC", 1.0, 990, 990)
        sheet = SheetType("S1", "大板", "SPCC", 1.0, 1000, 1000, edge_loss=10, max_sheets=1)
        result = optimize([product], [sheet], 1)[0]
        self.assertTrue(result.complete)
        self.assertEqual(result.shortage_by_product["P1"], 0)

    def test_rotation(self):
        product = ProductSpec("P1", "SPCC", 1.0, 900, 500, rotation_allowed=True)
        sheet = SheetType("S1", "大板", "SPCC", 1.0, 600, 1000, edge_loss=0)
        result = optimize([product], [sheet], 1)[0]
        self.assertTrue(result.complete)
        self.assertTrue(result.sheets[0].placements[0].rotated)

    def test_rotation_can_be_forbidden(self):
        product = ProductSpec("P1", "SPCC", 1.0, 900, 500, rotation_allowed=False)
        sheet = SheetType("S1", "大板", "SPCC", 1.0, 600, 1000, edge_loss=0, max_sheets=1)
        result = optimize([product], [sheet], 1)[0]
        self.assertFalse(result.complete)

    def test_multiple_products_full_yield(self):
        products = [
            ProductSpec("P1", "SPCC", 1.0, 500, 1000, required_qty=1),
            ProductSpec("P2", "SPCC", 1.0, 500, 500, required_qty=2),
        ]
        sheet = SheetType("S1", "大板", "SPCC", 1.0, 1000, 1000, edge_loss=0)
        results = optimize(products, [sheet], 2)
        complete = [result for result in results if result.complete]
        self.assertTrue(complete)
        self.assertAlmostEqual(complete[0].yield_rate, 100.0, places=6)
        self.assertEqual(validate_layout(complete[0]), [])

    def test_multiple_sheet_types_and_limits(self):
        products = [ProductSpec("P1", "SPCC", 1.0, 500, 500, required_qty=5)]
        sheets = [
            SheetType("S1", "大板A", "SPCC", 1.0, 1000, 1000, max_sheets=1, edge_loss=0),
            SheetType("S2", "大板B", "SPCC", 1.0, 500, 1000, max_sheets=2, edge_loss=0),
        ]
        result = optimize(products, sheets, 2)[0]
        self.assertTrue(result.complete)
        used = {plan.sheet_type.id for plan in result.sheets}
        self.assertIn("S1", used)
        self.assertIn("S2", used)
        self.assertEqual(validate_layout(result), [])

    def test_blank_sheet_quantity_calculates_required_sheet_count(self):
        product = ProductSpec("P1", "SPCC", 1.0, 500, 500, required_qty=9)
        sheet = SheetType("S1", "大板", "SPCC", 1.0, 1000, 1000, max_sheets=None, edge_loss=0)
        result = optimize([product], [sheet], 2)[0]
        self.assertTrue(result.complete)
        self.assertEqual(len(result.sheets), 3)

    def test_identical_sheet_patterns_are_grouped(self):
        sheet = SheetType("S1", "大板", "SPCC", 1.0, 1000, 1000, edge_loss=10)
        placements = [Placement("P1", 10, 10, 200, 300, False)]
        plans = [
            SheetPlan(sheet, 1, list(placements)),
            SheetPlan(sheet, 2, list(placements)),
            SheetPlan(sheet, 3, [Placement("P1", 10, 10, 300, 200, True)]),
        ]
        groups = group_sheet_plans(plans)
        self.assertEqual(len(groups), 2)
        self.assertEqual(groups[0][1], 2)
        self.assertEqual(groups[0][2], [1, 2])

    def test_square_sheet_whole_rotation_is_deduplicated(self):
        sheet = SheetType("S1", "大板", "SPCC", 1.0, 1219, 1219, edge_loss=10)
        vertical = SheetPlan(sheet, 1, [
            Placement("P1", index * 155, 0, 155, 1219, False)
            for index in range(7)
        ])
        horizontal = SheetPlan(sheet, 2, [
            Placement("P1", 0, index * 155, 1219, 155, True)
            for index in range(7)
        ])
        groups = group_sheet_plans([vertical, horizontal])
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0][1], 2)


class ExportTests(unittest.TestCase):
    def test_export_can_be_reopened(self):
        product = ProductSpec("P1", "SPCC", 1.0, 500, 500)
        sheet = SheetType("S1", "大板", "SPCC", 1.0, 1000, 1000, edge_loss=0)
        result = optimize([product], [sheet], 1)[0]
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "result.xlsx"
            export_result(path, result, [product], [sheet])
            wb = load_workbook(path, data_only=False)
            self.assertEqual(wb.sheetnames, ["計算条件", "推奨結果", "配置図"])
            self.assertEqual(wb["計算条件"]["B4"].value, "画面から手入力")
            self.assertEqual(wb["計算条件"]["F8"].value, "自動計算")
            self.assertEqual(wb["計算条件"]["I8"].value, 1)
            self.assertEqual(wb["推奨結果"]["A3"].value, "全数充足")
            self.assertEqual(wb["配置図"]["C3"].value, "P1")


if __name__ == "__main__":
    unittest.main()
