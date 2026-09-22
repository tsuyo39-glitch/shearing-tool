from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from openpyxl import load_workbook
from core.auto_exporter import export_auto
from core.auto_length import optimize_auto
from core.weight import plate_weight_kg
from tests.test_auto_length import product


class ExportTests(unittest.TestCase):
    def test_report_images_and_exact_quantities(self):
        products=[product()]
        result=optimize_auto(products)[0]
        with tempfile.TemporaryDirectory() as tmp:
            path=export_auto(Path(tmp)/"report.xlsx",result,products,(1219,10,2438,0),2)
            wb=load_workbook(path)
            self.assertEqual(wb.sheetnames,["結果一覧","図解一覧"])
            self.assertEqual(wb["結果一覧"]["D2"].value,2)
            self.assertEqual(wb["結果一覧"]["D3"].value,2)
            for name in wb.sheetnames[1:]:
                self.assertEqual(len(wb[name]._images),1)
                self.assertEqual(wb[name].page_setup.fitToWidth,1)
                self.assertEqual(wb[name].page_setup.fitToHeight,1)
                self.assertEqual(str(wb[name].page_setup.paperSize),str(wb[name].PAPERSIZE_A4))
                self.assertEqual(wb[name].page_setup.orientation,"landscape")
            with ZipFile(path) as archive:
                self.assertEqual(len([n for n in archive.namelist() if n.startswith("xl/media/")]),1)

    def test_gui_exports_selected_snapshot(self):
        import tkinter as tk
        from unittest.mock import patch
        from app_auto_length import AutoLengthApp
        root=tk.Tk()
        root.withdraw()
        try:
            app=AutoLengthApp(root)
            app.current_products=[product()]
            app.current_settings=(1219,10,2438,0)
            app.results=optimize_auto(app.current_products)
            app.result_tree.insert("","end",iid="1")
            app.result_tree.selection_set("1")
            app.settings[0].set("999")
            with patch("app_auto_length.filedialog.asksaveasfilename",return_value="sample.xlsx"),patch("app_auto_length.export_auto") as export,patch("app_auto_length.messagebox.showinfo"):
                app.export_excel()
                export.assert_not_called()
                app.result_valid=True
                app.result_revision=app.input_revision
                app.export_excel()
                self.assertIs(export.call_args.args[1],app.results[1])
                self.assertEqual(export.call_args.args[3],(1219,10,2438,0))
        finally:
            root.destroy()


def _table(ws, first_header):
    """見出し行の次から、A列が空になるまでの明細行を返す。"""
    rows=[row for row in ws.iter_rows(values_only=True)]
    start=next(i for i,row in enumerate(rows) if row[0]==first_header)
    body=[]
    for row in rows[start+1:]:
        if row[0] is None:
            break
        body.append(row)
    return body


class WeightTests(unittest.TestCase):
    def test_excel_weight_totals_match_rows(self):
        products=[product(300,400,40,thickness=1.6),product(250,900,20,id="P2",thickness=1.6)]
        result=optimize_auto(products)[0]
        self.assertTrue(result.weight_available)
        with tempfile.TemporaryDirectory() as tmp:
            ws=load_workbook(export_auto(Path(tmp)/"w.xlsx",result,products,(1219,10,2438,0,0),1))["結果一覧"]
            self.assertEqual([ws.cell(4,c).value for c in (1,3,5)],["大板重量(kg)","製品重量(kg)","端材重量(kg)"])
            sheet_kg,product_kg,scrap_kg=(ws.cell(4,c).value for c in (2,4,6))
            self.assertAlmostEqual(sheet_kg,result.sheet_weight,places=6)
            self.assertAlmostEqual(sheet_kg-product_kg,scrap_kg,places=6)
            sheet_rows=_table(ws,"規格")
            self.assertAlmostEqual(sum(row[7] for row in sheet_rows),sheet_kg,places=6)
            product_rows={row[0]:row for row in _table(ws,"製品ID")}
            self.assertEqual(set(product_rows),{"P1","P2"})
            for p in products:
                row=product_rows[p.id]
                unit=plate_weight_kg(p.spec,p.thickness,p.width,p.length)[0]
                self.assertAlmostEqual(row[12],unit,places=6)
                self.assertAlmostEqual(row[13],unit*row[9],places=6)
            self.assertAlmostEqual(sum(row[13] for row in product_rows.values()),product_kg,places=6)

    def test_excel_marks_weight_unknown_without_thickness(self):
        products=[product()]
        result=optimize_auto(products)[0]
        self.assertFalse(result.weight_available)
        with tempfile.TemporaryDirectory() as tmp:
            ws=load_workbook(export_auto(Path(tmp)/"w.xlsx",result,products,(1219,10,2438,0,0),1))["結果一覧"]
            self.assertEqual([ws.cell(4,c).value for c in (2,4,6)],["未計算"]*3)
            self.assertEqual(_table(ws,"規格")[0][7],"未計算")
            self.assertEqual(_table(ws,"製品ID")[0][12:14],("未計算","未計算"))
            self.assertIn("合計重量は未計算",ws.cell(8,2).value)

    def test_gui_candidate_list_and_detail_show_weight(self):
        import tkinter as tk
        from unittest.mock import patch
        import app_auto_length
        from app_auto_length import AutoLengthApp
        with tempfile.TemporaryDirectory() as tmp, patch.object(app_auto_length,"APP_DIR",Path(tmp)):
            root=tk.Tk()
            root.withdraw()
            try:
                app=AutoLengthApp(root)
                self.assertEqual(app.result_tree.heading("scrap")["text"],"端材kg")
                for products,expected in (([product(300,400,40,thickness=1.6)],True),([product()],False)):
                    app.current_products=products
                    app.calculation_revision=app.input_revision
                    app.pending=(optimize_auto(products),None)
                    app.result_tree.delete(*app.result_tree.get_children())
                    app._poll()
                    scrap=str(app.result_tree.item("0")["values"][3])
                    detail=app.detail.get("1.0","end")
                    if expected:
                        self.assertAlmostEqual(float(scrap.replace(",","")),app.results[0].scrap_weight,places=1)
                        self.assertIn("端材",detail)
                        self.assertIn("kg",detail)
                    else:
                        self.assertEqual(scrap,"未計算")
                        self.assertIn("重量：板厚が空欄のため未計算",detail)
            finally:
                root.destroy()
