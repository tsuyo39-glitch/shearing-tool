"""Excel report for the desktop application, including embedded layout images."""
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill
from .exporter import _title, _header, _finish, THIN
from .optimizer import group_sheet_plans, validate_layout
from .weight import coating_detail, plate_weight_kg

COLORS = ["#b9ddf3", "#bee5cb", "#ffe09d", "#d9cdf4", "#fac7b5", "#aadedd"]


def diagram(plan, colors):
    image = Image.new("RGB", (1200, 900), "white")
    d = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/meiryo.ttc")
    if not font_path.exists():
        raise ValueError("図解用の日本語フォント（メイリオ）が見つかりません。")
    font = ImageFont.truetype(str(font_path), 24)
    small = ImageFont.truetype(str(font_path), 18)
    s = plan.sheet_type
    scale = min(930/s.width, 640/s.length)
    x, y = 190, 100
    w, h = s.width*scale, s.length*scale
    d.text((x, 15), f"大板幅 {s.width:g} mm", fill="#164e63", font=font)
    d.line((x, 70, x+w, 70), fill="#334155", width=2)
    for xx in (x, x+w):
        d.line((xx, 60, xx, 80), fill="#334155", width=2)
    d.text((5, y+h/2-35), f"大板長さ\n{s.length:g} mm", fill="#164e63", font=font)
    d.rectangle((x,y,x+w,y+h), fill="#64748b", outline="#334155", width=3)
    e = s.edge_loss*scale
    le=s.length_loss*scale
    d.rectangle((x+e,y+le,x+w,y+h-le), fill="#e5e7eb")
    for p in plan.placements:
        bounds=(x+p.x*scale,y+p.y*scale,x+(p.x+p.width)*scale,y+(p.y+p.length)*scale)
        d.rectangle(bounds, fill=colors[p.product_id], outline="#334155", width=2)
        label=f"{p.product_id}\n{p.width:g}×{p.length:g}" + ("\n90°回転" if p.rotated else "")
        box=d.multiline_textbbox((0,0),label,font=small)
        if box[2]+8 < bounds[2]-bounds[0] and box[3]+8 < bounds[3]-bounds[1]:
            d.multiline_text(((bounds[0]+bounds[2]-box[2])/2,(bounds[1]+bounds[3]-box[3])/2),label,fill="#142b40",font=small,align="center")
        elif bounds[2]-bounds[0]>30 and bounds[3]-bounds[1]>25:
            d.text((bounds[0]+3,bounds[1]+3),p.product_id,fill="#142b40",font=small)
    d.text((30, 790), f"濃い灰色＝幅ロス（片側 {s.edge_loss:g} mm）／薄い灰色＝端材・切断代",fill="#334155",font=font)
    d.text((30, 835), f"色付き＝製品（下表のIDと対応）／製品間切断代 {s.cut_allowance:g} mm",fill="#334155",font=font)
    stream=BytesIO()
    image.save(stream,format="PNG")
    stream.seek(0)
    return stream


def export_auto(destination, result, products, settings, candidate_no=1):
    errors=validate_layout(result)
    if errors:
        raise ValueError("配置に問題があります："+" / ".join(errors[:3]))
    path=Path(destination)
    path.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook()
    ws=wb.active
    ws.title="結果一覧"
    _title(ws,"大板長さ自動計算 — 取り合わせ結果",14)
    ws.append(["出力日時",datetime.now().strftime("%Y-%m-%d %H:%M:%S"),"候補",candidate_no])
    ws.append(["状態","全数充足" if result.complete else "不足あり","大板枚数",len(result.sheets),"歩留り",result.yield_rate/100,"歩損",result.loss_rate/100])
    ws["F3"].number_format=ws["H3"].number_format="0.00%"
    totals=[result.sheet_weight,result.product_weight,result.scrap_weight] if result.weight_available else ["未計算"]*3
    ws.append(["大板重量(kg)",totals[0],"製品重量(kg)",totals[1],"端材重量(kg)",totals[2]])
    for column in "BDF":
        ws[f"{column}{ws.max_row}"].number_format="#,##0.0"
    ws.append(["大板幅(mm)",settings[0],"幅ロス(mm)",settings[1],"最大長さ(mm)",settings[2],"切断代(mm)",settings[3],
               "前後各ロス(mm)",settings[4] if len(settings)>4 else 0,"長さ丸め(mm)",settings[5] if len(settings)>5 else 0])
    merged=[]
    if not result.timed_out:
        note="探索候補（最適性の保証なし）"
    elif result.complete:
        note="必要数量を満たす候補。計算時間・中止で探索を打ち切ったため、時間を延ばすとより良い候補が出る場合があります。"
    else:
        note="時間制限・中止時点の暫定候補。必要数量を満たしていません。"
    ws.append(["注記",note])
    merged.append(ws.max_row)
    ws.append(["説明","計算時点の結果です。入力欄を変更した場合はアプリで再計算してください。"])
    merged.append(ws.max_row)
    if not result.weight_available:
        weight_note="板厚が空欄の製品があるため合計重量は未計算です。板厚が入力された行のみ重量を表示しています。"
    elif result.warnings:
        weight_note=" / ".join(result.warnings)
    else:
        weight_note="全製品の板厚が入力済みのため重量を算出しています。"
    ws.append(["重量の注記",weight_note])
    merged.append(ws.max_row)
    ws.append([])
    ws.append(["規格","板厚(mm)","めっき記号","めっき量定数 kg/m²","大板幅(mm)","自動長さ(mm)","必要大板枚数","大板重量(kg)"])
    _header(ws[ws.max_row])
    counts=Counter((s.sheet_type.spec,s.sheet_type.thickness,s.sheet_type.width,s.sheet_type.length,s.sheet_type.weight_spec) for s in result.sheets)
    for (spec,t,w,h,weight_spec),n in counts.items():
        code,entry=coating_detail(weight_spec)
        ws.append([spec or "未指定",t if t is not None else "未指定",code or "なし",
                   entry.constant_kg_m2 if entry else ("未登録" if code else "—"),w,h,n,
                   plate_weight_kg(weight_spec,t,w,h,n)[0] if t is not None else "未計算"])
        ws.cell(ws.max_row,8).number_format="#,##0.0"
    ws.append([])
    ws.append(["製品ID","製品名","規格","板厚(mm)","めっき記号","めっき量定数 kg/m²","幅(mm)","長さ(mm)",
               "必要枚数","配置枚数","不足枚数","回転許可","単重(kg)","配置重量(kg)"])
    _header(ws[ws.max_row])
    for p in products:
        placed=result.placed_by_product.get(p.id,0)
        code,entry=coating_detail(p.weight_spec)
        unit=plate_weight_kg(p.weight_spec,p.thickness,p.width,p.length)[0] if p.thickness is not None else None
        ws.append([p.id,p.name,p.spec or "未指定",p.thickness,code or "なし",
                   entry.constant_kg_m2 if entry else ("未登録" if code else "—"),
                   p.width,p.length,p.required_qty,placed,
                   result.shortage_by_product[p.id],"可" if p.rotation_allowed else "不可",
                   "未計算" if unit is None else unit,
                   "未計算" if unit is None else unit*placed])
        for column in (13,14):
            ws.cell(ws.max_row,column).number_format="#,##0.000"
    ws.append([])
    ws.append(["計算式","有効幅＝大板幅－幅ロス。大板長さ＝配置の最下端＋後端ロス（前端ロスは配置座標に含む）。"])
    merged.append(ws.max_row)
    ws.append(["歩留り","製品総面積 ÷ 投入大板総面積 × 100。歩損＝100－歩留り。"])
    merged.append(ws.max_row)
    ws.append(["重量","JIS G 3302-2010 表7：単位質量(kg/m²)＝表示厚さ(mm)×7.85＋めっき量定数(kg/m²)。単重＝単位質量×幅×長さ。端材重量＝大板重量－製品重量。"])
    merged.append(ws.max_row)
    ws.append(["重量の前提","表示厚さはめっき前の原板厚さ。めっき量定数はJIS G 3302-2010 表8の値（付着量g/m²とは別物）。板厚が空欄の製品が1つでもあると合計重量を未計算にします。"])
    merged.append(ws.max_row)
    colors={p.id:COLORS[i%len(COLORS)] for i,p in enumerate(products)}
    product_map={p.id:p for p in products}
    _finish(ws,[16,24,20,16,16,17,15,15,13,13,13,12,15,17])
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if cell.row>1 and cell.fill.fgColor.rgb!="001F4E78":
                    cell.font=Font(name="Yu Gothic",size=11)
                cell.alignment=Alignment(vertical="center",wrap_text=True)
                if cell.value is not None:
                    cell.border=Border(bottom=THIN)
                if isinstance(cell.value,str) and cell.value.startswith("="):
                    cell.data_type="s"
        if sheet.title=="結果一覧":
            for r in merged:
                sheet.merge_cells(start_row=r,start_column=2,end_row=r,end_column=14)
        else:
            for r in range(sheet.max_row-2,sheet.max_row+1):
                sheet.merge_cells(start_row=r,start_column=2,end_row=r,end_column=8)
        for row in sheet:
            if row[0].row<5 or sheet.title=="結果一覧" or row[0].row>=40:
                sheet.row_dimensions[row[0].row].height=34
    from .layout_overview import overview
    from openpyxl.worksheet.page import PageMargins
    layout=wb.create_sheet("図解一覧")
    picture=ExcelImage(overview(result,colors,diagram))
    picture.width,picture.height=1200,800
    layout.add_image(picture,"A1")
    for col in "ABCDEFGHIJKL":
        layout.column_dimensions[col].width=14
    for row in range(1,41):
        layout.row_dimensions[row].height=15
    layout.sheet_view.showGridLines=False
    layout.page_setup.orientation="landscape"
    layout.page_setup.paperSize=layout.PAPERSIZE_A4
    layout.page_setup.fitToWidth=1
    layout.page_setup.fitToHeight=1
    layout.sheet_properties.pageSetUpPr.fitToPage=True
    layout.page_margins=PageMargins(left=0.2,right=0.2,top=0.2,bottom=0.2,header=0,footer=0)
    layout.print_area="A1:L40"
    layout.print_options.horizontalCentered=True
    wb.save(path)
    return path
