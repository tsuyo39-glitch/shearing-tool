from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import OptimizationResult, ProductSpec, SheetType

NAVY = "1F4E78"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="A6A6A6")


def _title(ws, text: str, end_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, text)
    cell.font = Font(name="Yu Gothic", size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28


def _header(row) -> None:
    for cell in row:
        cell.font = Font(name="Yu Gothic", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN)


def _finish(ws, widths: list[float]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = ws.calculate_dimension()


def export_result(
    destination: str | Path,
    result: OptimizationResult,
    products: list[ProductSpec],
    sheet_types: list[SheetType],
) -> Path:
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    product_map = {product.id: product for product in products}

    wb = Workbook()
    conditions = wb.active
    conditions.title = "計算条件"
    _title(conditions, "シャーリング取り合わせ 計算条件", 9)
    conditions.append(["項目", "値"])
    _header(conditions[2])
    conditions.append(["計算日時", datetime.now()])
    conditions["B3"].number_format = "yyyy-mm-dd hh:mm:ss"
    conditions.append(["製品入力方式", "画面から手入力"])
    conditions.append(["計算状態", "暫定最良解" if result.timed_out else "探索完了"])
    conditions.append([])
    conditions.append(["大板名", "規格", "板厚(mm)", "幅(mm)", "長さ(mm)", "入力枚数", "外周ロス(mm/辺)", "切断代(mm)", "必要大板枚数"])
    _header(conditions[7])
    used = Counter(plan.sheet_type.id for plan in result.sheets)
    for sheet in sheet_types:
        conditions.append([
            sheet.name, sheet.spec, sheet.thickness, sheet.width, sheet.length,
            sheet.max_sheets if sheet.max_sheets is not None else "自動計算",
            sheet.edge_loss, sheet.cut_allowance, used.get(sheet.id, 0),
        ])
    _finish(conditions, [18, 22, 12, 12, 12, 14, 18, 14, 12])

    summary = wb.create_sheet("推奨結果")
    _title(summary, "推奨取り合わせ結果", 11)
    summary.append(["状態", "使用大板枚数", "歩留り率", "歩損率", "大板重量(kg)", "製品重量(kg)", "端材重量(kg)", "切断回数(概算)", "警告", "備考", "入力方式"])
    _header(summary[2])
    summary.append([
        "全数充足" if result.complete else "不足あり",
        len(result.sheets), result.yield_rate / 100, result.loss_rate / 100,
        result.sheet_weight, result.product_weight, result.scrap_weight,
        result.cut_count, " / ".join(result.warnings),
        "時間上限到達" if result.timed_out else "", "手入力",
    ])
    summary["C3"].number_format = summary["D3"].number_format = "0.00%"
    for cell in summary[3][4:7]:
        cell.number_format = "#,##0.0"
    summary.append([])
    summary.append(["製品ID", "製品名", "規格", "板厚(mm)", "幅(mm)", "長さ(mm)", "必要枚数", "配置枚数", "不足枚数", "回転許可", "備考"])
    _header(summary[5])
    for product in products:
        placed = result.placed_by_product.get(product.id, 0)
        summary.append([
            product.id, product.name, product.spec, product.thickness, product.width, product.length,
            product.required_qty, placed, max(0, product.required_qty - placed),
            "可" if product.rotation_allowed else "不可", product.remarks,
        ])
    _finish(summary, [12, 18, 22, 12, 12, 12, 12, 12, 12, 12, 22])

    layout = wb.create_sheet("配置図")
    _title(layout, "大板別 配置図・座標", 11)
    layout.append(["大板名", "使用板No", "製品ID", "X(mm)", "Y(mm)", "配置幅(mm)", "配置長さ(mm)", "元幅(mm)", "元長さ(mm)", "回転", "切断順"])
    _header(layout[2])
    row_no = 3
    for plan in result.sheets:
        for cut_order, placement in enumerate(plan.placements, 1):
            product = product_map[placement.product_id]
            layout.append([
                plan.sheet_type.name, plan.sheet_no, placement.product_id,
                placement.x, placement.y, placement.width, placement.length,
                product.width, product.length, "90度" if placement.rotated else "なし", cut_order,
            ])
            fill = GREEN if cut_order % 2 else BLUE
            for cell in layout[row_no]:
                cell.fill = PatternFill("solid", fgColor=fill)
            row_no += 1
    _finish(layout, [18, 12, 12, 12, 12, 14, 14, 12, 12, 10, 10])

    wb.save(path)
    return path


def default_output_path(base_dir: str | Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path(base_dir) / "計算結果" / f"シャーリング取り合わせ_{timestamp}.xlsx"
